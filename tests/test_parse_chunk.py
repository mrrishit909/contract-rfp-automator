from __future__ import annotations

import io

import pdfplumber
import pytest
from docx import Document

from app.chunk import chunk, definitions
from app.parse import Line, UnsupportedFile, ocr_available, parse


@pytest.fixture(scope="module")
def lines(pdf: bytes) -> list[Line]:
    return parse("tsa.pdf", pdf)


def test_pdf_pages_text_and_running_lines(lines: list[Line]) -> None:
    assert max(ln.page for ln in lines) == 52
    assert not any(ln.text.startswith("Source: REYNOLDS") for ln in lines)       # the footer printed on every page
    assert not any(ln.text.strip().isdigit() for ln in lines)                    # bare page numbers
    assert any("TRANSITION SERVICES AGREEMENT" in ln.text for ln in lines if ln.page == 1)


def test_clauses_follow_the_contract_s_own_structure(lines: list[Line]) -> None:
    clauses = chunk(lines)
    by_number = {c.number: c for c in clauses if c.page <= 27}
    assert [by_number[n].title for n in ("5.4", "6.2", "9.1", "9.4", "10.8", "10.10")] == [
        "Payment", "Termination of Services", "Indemnification", "Limitations", "Governing Law, etc", "Assignment"]
    assert by_number["9.4"].section == "Article IX REMEDIES" and by_number["5.4"].section == "Article V FEES"
    assert "WILLFUL MISCONDUCT" in by_number["9.4"].text and "last business day of the month" in by_number["5.4"].text
    sections = list(dict.fromkeys(c.section for c in clauses if c.section))
    assert [s.split()[1] for s in sections[:10]] == ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
    assert all(len(c.text) <= 6100 for c in clauses) and all(c.text.strip() for c in clauses)
    kept = sum(len(c.text) for c in clauses) / sum(len(ln.text) for ln in lines)
    assert kept > 0.95                                                           # chunking loses almost no text
    terms = definitions(clauses)
    assert len(terms) > 20 and terms["Term"].startswith("“Term” has the meaning")


def L(*texts: str, page: int = 1) -> list[Line]:
    return [Line(t, page) for t in texts]


def test_table_of_contents_addresses_and_wrapped_lines_are_not_clauses() -> None:
    body = "The Supplier shall deliver the Products to the Customer at the agreed place and on the agreed delivery date."
    lines = L("MASTER SUPPLY AGREEMENT", "TABLE OF CONTENTS",
              "1 Definitions 3", "2 Supply of Products 4", "3 Fees 7", "4 Term 9", "5 General 11",
              "1. DEFINITIONS", "1.1 “Product” means the goods listed in Schedule A.",
              "2. SUPPLY OF PRODUCTS", f"2.1 Delivery. {body}",
              "(a) Risk passes on delivery.", "(b) Title passes on payment.",
              "Notices shall be sent to the address below, marked for the attention of the General Counsel at",
              "75 Park Plaza, Boston, MA 02116 USA.", "30 days after the date of this Agreement the fees are due.",
              "Termination for Convenience", f"Either party may terminate on notice. {body}")
    clauses = chunk(lines)
    assert [(c.number, c.title) for c in clauses] == [
        ("", "Preamble"), ("1.1", "Section 1.1"), ("2.1", "Delivery"), ("", "Termination for Convenience")]
    assert clauses[1].text == "“Product” means the goods listed in Schedule A." and clauses[1].section == "1 DEFINITIONS"
    delivery = next(c for c in clauses if c.number == "2.1")
    assert "75 Park Plaza" in delivery.text and "30 days after" in delivery.text
    assert delivery.text.count("\n") >= 2 and delivery.section == "2 SUPPLY OF PRODUCTS"      # (a), (b) start paragraphs
    assert not any("Definitions 3" in c.text or "Fees 7" in c.text for c in clauses)          # the TOC is gone


def test_article_title_on_the_next_line_and_long_clause_split() -> None:
    para = "The Provider shall perform the Services with reasonable skill and care. " * 12
    lines = L("ARTICLE IV", "LIMITATION OF LIABILITY", "Section 4.1 Cap. Liability is capped.",
              "Section 4.2 Details. " + para, *[f"({chr(97 + i)}) {para}" for i in range(9)])
    clauses = chunk(lines)
    assert clauses[0].number == "4.1" and clauses[0].section == "Article IV LIMITATION OF LIABILITY"
    assert [c.title for c in clauses[1:]] == ["Details (part 1)", "Details (part 2)"]


def test_docx_headings_and_tables() -> None:
    d = Document()
    d.add_heading("Payment Terms", level=1)
    d.add_paragraph("Customer shall pay each invoice within sixty (60) days after receipt of a correct invoice.")
    table = d.add_table(rows=2, cols=2)
    for cell, text in zip([c for r in table.rows for c in r.cells], ["Service", "Fee", "Hosting", "USD 500"]):
        cell.text = text
    d.add_heading("Governing Law", level=1)
    d.add_paragraph("This Agreement is governed by the laws of the State of Delaware, without regard to conflict rules.")
    buf = io.BytesIO()
    d.save(buf)
    lines = parse("terms.docx", buf.getvalue())
    assert [ln.kind for ln in lines] == ["heading", "text", "table", "table", "heading", "text"]
    clauses = chunk(lines)
    assert [c.title for c in clauses] == ["Payment Terms", "Governing Law"] and "Hosting | USD 500" in clauses[0].text


def test_unsupported_and_empty_files() -> None:
    with pytest.raises(UnsupportedFile):
        parse("notes.txt", b"hello")
    with pytest.raises(UnsupportedFile, match="no text"):
        parse("empty.docx", (lambda b: (Document().save(b), b.getvalue())[1])(io.BytesIO()))


@pytest.mark.skipif(not ocr_available(), reason="Tesseract not installed")
def test_scanned_image_and_scanned_pdf_go_through_ocr(pdf: bytes, lines: list[Line]) -> None:
    with pdfplumber.open(io.BytesIO(pdf)) as doc:
        scan = doc.pages[9].to_image(resolution=200).original.convert("RGB")     # page 10 as a picture: no text layer
    png, as_pdf = io.BytesIO(), io.BytesIO()
    scan.save(png, "PNG")
    scan.save(as_pdf, "PDF")
    native = {w for ln in lines if ln.page == 10 for w in ln.text.lower().split()}
    for name, data in (("scan.png", png.getvalue()), ("scan.pdf", as_pdf.getvalue())):
        words = {w for ln in parse(name, data) for w in ln.text.lower().split()}
        assert len(words & native) / len(native) > 0.9
    assert any(c.number == "5.4" for c in chunk(parse("scan.png", png.getvalue())))  # clause structure survives OCR
