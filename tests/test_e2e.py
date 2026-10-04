"""End to end through the HTTP API: upload the 52-page agreement, wait for the background analysis, use the result."""
from __future__ import annotations

import time
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from app import api
from app.analyze import Review, analyze
from app.index import HybridIndex
from conftest import FakeLLM

AUTH = {"Authorization": "Bearer test-token"}


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, llm: FakeLLM, index: HybridIndex) -> Any:
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    api.configure(llm, index)
    with TestClient(api.app) as c:          # the context manager keeps one event loop alive for the background task
        yield c


def wait_ready(client: TestClient, doc_id: str, seconds: float = 120) -> dict[str, Any]:
    deadline = time.time() + seconds
    while time.time() < deadline:
        doc: dict[str, Any] = client.get(f"/api/documents/{doc_id}", headers=AUTH).json()
        if doc["status"] != "processing":
            return doc
        time.sleep(0.2)
    raise AssertionError("analysis did not finish")


def test_fifty_page_document_workflow(client: TestClient, pdf: bytes, llm: FakeLLM) -> None:
    started = time.time()
    r = client.post("/api/documents", headers=AUTH, files={"file": ("tsa.pdf", pdf)}, data={"perspective": "the Company (RCP)"})
    assert r.status_code == 202 and r.json()["status"] == "processing"
    doc_id = r.json()["id"]
    doc = wait_ready(client, doc_id)
    assert doc["status"] == "ready" and doc["pages"] == 52 and time.time() - started < 120
    clauses = doc["clauses"]
    assert len(clauses) >= 60 and all(c["review"]["risk"] in ("Low", "Medium", "High") for c in clauses)
    card = doc["scorecard"]
    assert card["High"] + card["Medium"] + card["Low"] == len(clauses) and card["High"] >= 1 and card["overall"] == "High"
    assert doc["perspective"] == "the Company (RCP)" and "the Company (RCP)" in llm.chats[0][0]["content"]

    # every risky clause has a redraft and a diff that reconstructs both sides
    risky = [c for c in clauses if c["review"]["risk"] != "Low"]
    for c in risky:
        ops = c["review"]["diff"]
        assert " ".join(o["a"] for o in ops if o["a"]).split() == c["text"].split()
        assert " ".join(o["b"] for o in ops if o["b"]).split() == c["review"]["suggested_text"].split()
    # defined terms used in a clause were looked up elsewhere in the document and sent along
    assert any('"defined_terms": {"' in chat[-1]["content"] for chat in llm.chats)

    # apply one AI draft, export, undo
    target = next(c for c in risky if c["review"]["risk"] == "High")
    url = f"/api/documents/{doc_id}/clauses/{target['id']}/apply"
    assert client.post(url, headers=AUTH, json={"applied": True}).json() == {"id": target["id"], "applied": True}
    exported = client.get(f"/api/documents/{doc_id}/export", headers=AUTH).text
    assert target["review"]["suggested_text"] in exported and target["text"] not in exported
    client.post(url, headers=AUTH, json={"applied": False})
    assert target["text"] in client.get(f"/api/documents/{doc_id}/export", headers=AUTH).text
    low = next(c for c in clauses if c["review"]["risk"] == "Low")
    assert client.post(f"/api/documents/{doc_id}/clauses/{low['id']}/apply", headers=AUTH, json={"applied": True}).status_code == 409

    # hybrid search inside the document
    hits = client.get(f"/api/documents/{doc_id}/search", headers=AUTH, params={"q": "governing law jurisdiction courts"}).json()
    titles = {c["id"]: c["title"] for c in clauses}
    assert any("Governing Law" in titles[h["id"]] for h in hits[:3])
    listing = client.get("/api/documents", headers=AUTH).json()
    assert [(d["id"], d["status"], d["pages"]) for d in listing] == [(doc_id, "ready", 52)]


def test_upload_rejections_and_auth(client: TestClient) -> None:
    assert client.get("/api/documents").status_code == 401
    assert client.post("/api/documents", headers=AUTH, files={"file": ("evil.exe", b"MZ")}).status_code == 415
    assert client.get("/api/documents/nope", headers=AUTH).status_code == 404
    assert client.get("/api/documents/" + "0" * 32, headers=AUTH).status_code == 404
    r = client.post("/api/documents", headers=AUTH, files={"file": ("broken.pdf", b"this is not a pdf")})
    doc = wait_ready(client, r.json()["id"])
    assert doc["status"] == "failed" and doc["error"]
    assert client.get(f"/api/documents/{doc['id']}/export", headers=AUTH).status_code == 409


async def test_invalid_model_output_is_retried_then_flagged_for_a_human(index: HybridIndex, pdf: bytes) -> None:
    class Stubborn(FakeLLM):
        async def chat_json(self, messages: list[dict[str, str]], schema: dict[str, Any], name: str) -> dict[str, Any]:
            self.chats.append(messages)
            return {"clause_type": "x", "rule_id": "made_up_rule", "risk": "High", "explanation": "no redraft given", "suggested_text": None}

    llm = Stubborn()
    index.llm = llm
    d = Document_with_one_clause()
    out = await analyze("d" * 32, "one.docx", d, llm, index)
    review = out["clauses"][0]["review"]
    assert review["failed"] and review["risk"] == "Medium" and len(llm.chats) == 2          # tried twice, then handed to a person
    with pytest.raises(ValueError):
        Review(clause_type="x", rule_id=None, risk="High", explanation="e", suggested_text=" ")
    assert Review(clause_type="x", rule_id=None, risk="Low", explanation="e", suggested_text="ignored").suggested_text is None


def Document_with_one_clause() -> bytes:
    import io

    from docx import Document
    d = Document()
    d.add_heading("Limitation of Liability", level=1)
    d.add_paragraph("The Provider is liable only for losses caused by its gross negligence or wilful misconduct, and in no other case whatsoever.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
